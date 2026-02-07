#!/usr/bin/env python3
"""
Atualiza planilha com dados da análise visual Mistral
"""

import json
import re
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def extract_info_from_description(description_text: str) -> dict:
    """Extrai informações da descrição textual"""
    info = {
        "quartos": "",
        "banheiros": "1",
        "suites": "",
        "vagas": "",
        "area_terreno": "",
        "area_construida": "",
    }

    if not description_text:
        return info

    text_lower = description_text.lower()

    # Extrair quartos
    match = re.search(r"(\d+)\s*dormitório", text_lower)
    if match:
        info["quartos"] = match.group(1)

    # Extrair suítes
    match = re.search(r"(\d+)\s*suíte", text_lower)
    if match:
        info["suites"] = match.group(1)

    # Extrair vagas
    match = re.search(r"(\d+)\s*vaga", text_lower)
    if match:
        info["vagas"] = match.group(1)

    return info


def update_spreadsheet_with_mistral(
    batch_dir: str = ".tmp/batch_resultado",
    excel_file: str = ".tmp/planilha_10_imoveis.xlsx",
):
    """
    Atualiza planilha com dados da análise Mistral
    """

    batch_path = Path(batch_dir)

    # Encontrar todos os diretórios de imóveis
    property_dirs = sorted(
        [
            d
            for d in batch_path.iterdir()
            if d.is_dir() and d.name.startswith("property_")
        ]
    )

    print(
        f"📊 Atualizando planilha com dados Mistral para {len(property_dirs)} imóveis..."
    )

    # Carregar workbook existente ou criar novo
    if Path(excel_file).exists():
        wb = load_workbook(excel_file)
        ws = wb.active
    else:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Imóveis"

        # Headers
        headers = [
            "Endereço",
            "Bairro",
            "Telefone do informante",
            "Área do Terreno",
            "Área Construída",
            "Padrão construtivo",
            "Quartos + Wc",
            "Estado de conservação",
            "Indice fiscal",
            "Pavimentação",
            "Idade aparente",
            "Suítes",
            "Vagas de garagem coberta",
            "Data do evento",
            "Evento",
            "Geminada",
            "Multi",
            "Renda Média Bairro",
            "Valor total",
        ]

        header_fill = PatternFill(
            start_color="366092", end_color="366092", fill_type="solid"
        )
        header_font = Font(bold=True, color="FFFFFF", size=11)

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )

    # Limpar dados antigos (exceto header)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.value = None

    # Processar cada imóvel
    row = 2
    for prop_dir in property_dirs:
        json_file = prop_dir / "data.json"
        if not json_file.exists():
            continue

        with open(json_file, "r", encoding="utf-8") as f:
            data = json.load(f)

        # Extrair informações
        desc_text = data.get("description", {}).get("text", "")
        parsed = extract_info_from_description(desc_text)

        # Usar análise visual do Mistral se disponível
        visual = data.get("visual_analysis_mistral", {})
        if not visual:
            visual = data.get("visual_analysis", {})  # Fallback

        # Mapear campos
        endereco = data.get("basic_info", {}).get("title", "N/A")
        bairro = data.get("location", {}).get("neighborhood", "N/A")
        telefone = data.get("advertiser", {}).get("phone", "")
        area_terreno = parsed.get("area_terreno", "")
        area_construida = parsed.get("area_construida", "")
        padrao = visual.get("padrao_construtivo", "Médio").capitalize()
        quartos_wc = f"{parsed.get('quartos', '')}Q + {parsed.get('banheiros', '1')}B"
        conservacao = visual.get("estado_conservacao", "Regular").capitalize()
        indice_fiscal = ""  # Não disponível
        pavimentacao = visual.get("pavimentacao", "Indeterminado").capitalize()
        idade = visual.get("idade_aparente_anos", "")
        if idade:
            idade = f"{idade} anos"
        suites = parsed.get("suites", "")
        vagas = parsed.get("vagas", "")
        data_evento = datetime.now().strftime("%d/%m/%Y")
        evento = "Venda"
        geminada = "Não"
        multi = ""
        renda_bairro = ""
        valor = data.get("prices", {}).get("price_formatted", "N/A")

        values = [
            endereco,
            bairro,
            telefone,
            area_terreno,
            area_construida,
            padrao,
            quartos_wc,
            conservacao,
            indice_fiscal,
            pavimentacao,
            idade,
            suites,
            vagas,
            data_evento,
            evento,
            geminada,
            multi,
            renda_bairro,
            valor,
        ]

        # Escrever linha
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        row += 1

    # Ajustar largura das colunas
    column_widths = {
        "A": 45,  # Endereço
        "B": 22,  # Bairro
        "C": 20,  # Telefone
        "D": 15,  # Área Terreno
        "E": 15,  # Área Construída
        "F": 15,  # Padrão
        "G": 12,  # Quartos + WC
        "H": 18,  # Conservação
        "I": 12,  # Indice fiscal
        "J": 15,  # Pavimentação
        "K": 12,  # Idade
        "L": 10,  # Suítes
        "M": 18,  # Vagas
        "N": 12,  # Data
        "O": 12,  # Evento
        "P": 12,  # Geminada
        "Q": 10,  # Multi
        "R": 18,  # Renda
        "S": 20,  # Valor
    }

    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Congelar primeira linha
    ws.freeze_panes = "A2"

    # Salvar
    wb.save(excel_file)
    print(f"✅ Planilha atualizada: {excel_file}")
    print(f"📊 Total de imóveis: {row - 2}")
    print(f"🤖 Dados da análise: Mistral Pixtral")

    return excel_file


if __name__ == "__main__":
    import sys

    batch_dir = sys.argv[1] if len(sys.argv) > 1 else ".tmp/batch_resultado"
    excel_file = sys.argv[2] if len(sys.argv) > 2 else ".tmp/planilha_10_imoveis.xlsx"

    update_spreadsheet_with_mistral(batch_dir, excel_file)
