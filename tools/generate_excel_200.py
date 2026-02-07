#!/usr/bin/env python3
"""
Gera planilha final de 200 imóveis
"""

import json
import re
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def extract_info_from_description(description_text: str) -> dict:
    info = {"quartos": "", "banheiros": "1", "suites": "", "vagas": ""}
    if not description_text:
        return info
    text_lower = description_text.lower()
    match = re.search(r"(\d+)\s*dormitório", text_lower)
    if match:
        info["quartos"] = match.group(1)
    match = re.search(r"(\d+)\s*suíte", text_lower)
    if match:
        info["suites"] = match.group(1)
    match = re.search(r"(\d+)\s*vaga", text_lower)
    if match:
        info["vagas"] = match.group(1)
    return info


def generate_excel_200(
    batch_dir: str = ".tmp/batch_200_imoveis",
    excel_file: str = "planilha_200_imoveis_v2.xlsx",
):
    batch_path = Path(batch_dir)
    property_dirs = sorted(
        [
            d
            for d in batch_path.iterdir()
            if d.is_dir() and d.name.startswith("property_")
        ]
    )

    print(f"📊 Gerando planilha para {len(property_dirs)} imóveis...")

    wb = Workbook()
    ws = wb.active
    ws.title = "Imóveis"

    headers = [
        "Endereço",
        "Bairro",
        "Telefone",
        "Área Terreno",
        "Área Construída",
        "Padrão",
        "Quartos+WC",
        "Conservação",
        "Índice Fiscal",
        "Pavimentação",
        "Idade",
        "Suítes",
        "Vagas",
        "Data",
        "Evento",
        "Geminada",
        "Multi",
        "Renda Bairro",
        "Valor",
    ]

    header_fill = PatternFill(
        start_color="366092", end_color="366092", fill_type="solid"
    )
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = thin_border

    row = 2
    for prop_dir in property_dirs:
        json_file = prop_dir / "data.json"
        if not json_file.exists():
            continue

        with open(json_file, "r", encoding="utf-8") as f:
            data = json.load(f)

        desc_text = data.get("description", {}).get("text", "")
        parsed = extract_info_from_description(desc_text)
        visual = data.get("visual_analysis_mistral", {})

        endereco = data.get("basic_info", {}).get("title", "N/A")
        bairro = data.get("location", {}).get("neighborhood", "N/A")
        telefone = data.get("advertiser", {}).get("phone", "")
        area_terreno = ""
        area_construida = ""
        padrao = (visual.get("padrao_construtivo") or "Médio").capitalize()
        quartos_wc = f"{parsed.get('quartos', '')}Q+{parsed.get('banheiros', '1')}B"
        conservacao = (visual.get("estado_conservacao") or "Regular").capitalize()
        indice_fiscal = ""
        pavimentacao = (visual.get("pavimentacao") or "Indeterminado").capitalize()
        idade = visual.get("idade_aparente_anos") or ""
        if idade:
            idade = f"{idade} anos"
        suites = parsed.get("suites", "")
        vagas = parsed.get("vagas", "")
        data_evento = datetime.now().strftime("%d/%m/%Y")
        evento = "Venda"
        geminada = "Não"
        multi = ""
        renda_bairro = ""
        valor = data.get("prices", {}).get("price_formatted", "N/A")

        values = [
            endereco,
            bairro,
            telefone,
            area_terreno,
            area_construida,
            padrao,
            quartos_wc,
            conservacao,
            indice_fiscal,
            pavimentacao,
            idade,
            suites,
            vagas,
            data_evento,
            evento,
            geminada,
            multi,
            renda_bairro,
            valor,
        ]

        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        row += 1

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col)].width = 15
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 20

    ws.freeze_panes = "A2"
    wb.save(excel_file)

    print(f"✅ Planilha gerada: {excel_file}")
    print(f"📊 Total: {row - 2} imóveis")
    return excel_file


if __name__ == "__main__":
    import sys

    batch_dir = sys.argv[1] if len(sys.argv) > 1 else ".tmp/batch_200_imoveis"
    excel_file = sys.argv[2] if len(sys.argv) > 2 else "planilha_200_imoveis_v2.xlsx"
    generate_excel_200(batch_dir, excel_file)
