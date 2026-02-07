#!/usr/bin/env python3
"""
Batch Scraper v2 - Processa múltiplos imóveis
=============================================
Versão integrada que faz scraping + análise + export
"""

import asyncio
import json
import hashlib
import sys
import os
import re
from pathlib import Path
from datetime import datetime
from urllib.parse import urljoin

from playwright.async_api import async_playwright

# Config
BASE_URL = "https://www.infoimoveis.com.br"


async def scrape_single_property(url: str, output_dir: Path) -> dict:
    """
    Faz scraping de um único imóvel

    Returns:
        dict com dados do imóvel ou None se falhar
    """

    images_dir = output_dir / "images"
    images_dir.mkdir(parents=True, exist_ok=True)

    async with async_playwright() as p:
        browser = await p.chromium.launch(
            headless=False,  # CRITICAL: False = funciona
            args=[
                "--disable-blink-features=AutomationControlled",
                "--no-sandbox",
                "--disable-dev-shm-usage",
            ],
        )

        context = await browser.new_context(
            viewport={"width": 1920, "height": 1080},
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
        )

        page = await context.new_page()

        # Anti-detecção
        await page.add_init_script("""
            Object.defineProperty(navigator, 'webdriver', {
                get: () => undefined
            });
        """)

        print(f"  🌐 Navegando...")
        await page.goto(url, wait_until="domcontentloaded", timeout=60000)
        await page.wait_for_timeout(10000)  # Cloudflare

        # Verificar bloqueio
        content = await page.content()
        if "cloudflare" in content.lower() or "access denied" in content.lower():
            print(f"  ❌ BLOQUEADO!")
            await browser.close()
            return None

        print(f"  ✅ Cloudflare OK")

        # Inicializar dados
        data = {
            "metadata": {
                "source_url": url,
                "scraped_at": datetime.now().isoformat(),
            },
            "basic_info": {},
            "location": {},
            "areas": {},
            "rooms": {},
            "prices": {},
            "features": [],
            "description": {},
            "advertiser": {},
            "images": [],
        }

        # Título
        try:
            title = await page.locator(
                "h1.titulo, h1.title, .property-title h1"
            ).first.inner_text()
            data["basic_info"]["title"] = title.strip()
        except:
            pass

        # Preço - múltiplas tentativas
        try:
            # Tentar hidden input primeiro
            price = await page.input_value("#priceImovel")
            if price:
                data["prices"]["price_raw"] = price
                data["prices"]["price_formatted"] = f"R$ {float(price):,.0f}"
        except:
            try:
                price_text = await page.locator(
                    ".preco, .price, .valor"
                ).first.inner_text()
                data["prices"]["price_formatted"] = price_text.strip()
            except:
                pass

        # Localização
        try:
            neighborhood = await page.locator(
                ".bairro, .neighborhood"
            ).first.inner_text()
            data["location"]["neighborhood"] = neighborhood.strip()
        except:
            pass

        try:
            city = await page.locator(".cidade, .city").first.inner_text()
            data["location"]["city"] = city.strip()
        except:
            data["location"]["city"] = "Campo Grande"

        # Áreas
        try:
            area_total = await page.locator(
                ".area-terreno, .total-area"
            ).first.inner_text()
            data["areas"]["total"] = area_total.strip()
        except:
            pass

        try:
            area_construida = await page.locator(
                ".area-construida, .built-area"
            ).first.inner_text()
            data["areas"]["built"] = area_construida.strip()
        except:
            pass

        # Características (features)
        try:
            features = await page.locator(
                ".caracteristica, .feature, .item-caracteristica"
            ).all()
            for feat in features:
                try:
                    text = await feat.inner_text()
                    if text.strip():
                        data["features"].append(text.strip())
                except:
                    pass
        except:
            pass

        # Parse quartos/banheiros/vagas das features
        for feature in data["features"]:
            feature_lower = feature.lower()

            if "quarto" in feature_lower and not data["rooms"].get("bedrooms"):
                match = re.search(r"(\d+)", feature)
                if match:
                    data["rooms"]["bedrooms"] = int(match.group(1))

            if "suíte" in feature_lower or "suite" in feature_lower:
                match = re.search(r"(\d+)", feature)
                if match:
                    data["rooms"]["suites"] = int(match.group(1))

            if "wc" in feature_lower or "banheiro" in feature_lower:
                match = re.search(r"(\d+)", feature)
                if match:
                    data["rooms"]["bathrooms"] = int(match.group(1))

            if "vaga" in feature_lower:
                match = re.search(r"(\d+)", feature)
                if match:
                    data["rooms"]["parking_spaces"] = int(match.group(1))

        # Anunciante
        try:
            advertiser = await page.locator(
                ".anunciante, .advertiser-name"
            ).first.inner_text()
            data["advertiser"]["name"] = advertiser.strip()
        except:
            pass

        try:
            phone = await page.locator(".telefone, .phone").first.inner_text()
            data["advertiser"]["phone"] = phone.strip()
        except:
            pass

        # Descrição
        try:
            desc = await page.locator(".descricao, .description").first.inner_text()
            data["description"]["text"] = desc.strip()
        except:
            pass

        # Imagens
        image_urls = set()
        selectors = [
            'img[src*="/fotos/"]',
            'img[src*="stored/imoveis"]',
            'img[data-src*="/fotos/"]',
        ]

        for selector in selectors:
            try:
                imgs = await page.locator(selector).all()
                for img in imgs:
                    src = await img.get_attribute("src") or await img.get_attribute(
                        "data-src"
                    )
                    if src and ("foto" in src or "stored" in src):
                        if src.startswith("/"):
                            src = urljoin(BASE_URL, src)
                        image_urls.add(src)
            except:
                pass

        # Download imagens
        for idx, img_url in enumerate(sorted(image_urls), start=1):
            try:
                response = await page.request.get(img_url)
                if response.status == 200:
                    content = await response.body()
                    ext = ".jpg"
                    hash_str = hashlib.md5(img_url.encode()).hexdigest()[:8]
                    filename = f"image_{idx:02d}_{hash_str}{ext}"
                    filepath = images_dir / filename

                    with open(filepath, "wb") as f:
                        f.write(content)

                    data["images"].append(
                        {
                            "index": idx,
                            "url": img_url,
                            "local_path": str(filepath),
                            "filename": filename,
                        }
                    )
            except:
                pass

        # Salvar JSON
        json_file = output_dir / "data.json"
        with open(json_file, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

        await browser.close()

        return data


def analyze_with_groq(image_paths: list) -> dict:
    """
    Analisa imagens com Groq Vision
    """
    try:
        from groq import Groq
        import base64

        api_key = os.getenv("GROQ_API_KEY")
        if not api_key:
            print(f"    ⚠️ GROQ_API_KEY não encontrada")
            return {}

        client = Groq(api_key=api_key)

        # Analisar apenas as 3 primeiras imagens
        analyses = []
        for img_path in image_paths[:3]:
            try:
                with open(img_path, "rb") as f:
                    image_data = base64.b64encode(f.read()).decode("utf-8")

                prompt = """Analise esta foto de imóvel e retorne APENAS um JSON:
{
  "estado_conservacao": "novo|bom|regular|ruim",
  "idade_aparente_anos": <número ou null>,
  "padrao_construtivo": "alto|médio|baixo",
  "pavimentacao": "ceramica|porcelanato|madeira|cimento|indeterminado",
  "observacoes": "<breve descrição>"
}
Retorne APENAS o JSON, sem explicações."""

                response = client.chat.completions.create(
                    model="llama-3.2-90b-vision-preview",
                    messages=[
                        {
                            "role": "user",
                            "content": [
                                {"type": "text", "text": prompt},
                                {
                                    "type": "image_url",
                                    "image_url": {
                                        "url": f"data:image/jpeg;base64,{image_data}"
                                    },
                                },
                            ],
                        }
                    ],
                    temperature=0.1,
                    max_tokens=300,
                )

                result = response.choices[0].message.content.strip()

                # Extrair JSON
                if "```json" in result:
                    result = result.split("```json")[1].split("```")[0]
                elif "```" in result:
                    result = result.split("```")[1]

                analysis = json.loads(result.strip())
                analyses.append(analysis)

            except Exception as e:
                print(f"    ⚠️ Erro ao analizar imagem: {e}")
                continue

        # Agregar análises
        if not analyses:
            return {}

        # Votação para campos categóricos
        from collections import Counter

        conservation = Counter(
            [a.get("estado_conservacao", "indeterminado") for a in analyses]
        ).most_common(1)[0][0]
        padrao = Counter(
            [a.get("padrao_construtivo", "indeterminado") for a in analyses]
        ).most_common(1)[0][0]
        pavimentacao = Counter(
            [a.get("pavimentacao", "indeterminado") for a in analyses]
        ).most_common(1)[0][0]

        # Média para idade
        idades = [
            a.get("idade_aparente_anos")
            for a in analyses
            if a.get("idade_aparente_anos")
        ]
        idade_avg = sum(idades) / len(idades) if idades else None

        return {
            "estado_conservacao": conservation,
            "idade_aparente_anos": round(idade_avg) if idade_avg else None,
            "padrao_construtivo": padrao,
            "pavimentacao": pavimentacao,
            "analises_individuais": len(analyses),
        }

    except Exception as e:
        print(f"    ⚠️ Erro Groq: {e}")
        return {}


def export_to_excel_format(data: dict, output_file: str):
    """
    Exporta dados para Excel no formato da planilha do usuário
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = Workbook()
        ws = wb.active
        ws.title = "Imóveis"

        # Headers conforme planilha do usuário
        headers = [
            "Endereço",
            "Bairro",
            "Telefone do informante",
            "Área do Terreno",
            "Área Construída",
            "Padrão construtivo",
            "Quartos + Wc",
            "Estado de conservação",
            "Indice fiscal",
            "Pavimentação",
            "Idade aparente",
            "Suítes",
            "Vagas de garagem coberta",
            "Data do evento",
            "Evento",
            "Geminada",
            "Multi",
            "Renda Média Bairro",
            "Valor total",
        ]

        # Estilo header
        header_fill = PatternFill(
            start_color="366092", end_color="366092", fill_type="solid"
        )
        header_font = Font(bold=True, color="FFFFFF")

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Dados
        row = 2

        # Mapear dados
        endereco = data.get("basic_info", {}).get("title", "")
        bairro = data.get("location", {}).get("neighborhood", "")
        telefone = data.get("advertiser", {}).get("phone", "")
        area_terreno = data.get("areas", {}).get("total", "")
        area_construida = data.get("areas", {}).get("built", "")
        padrao = data.get("visual_analysis", {}).get("padrao_construtivo", "")
        quartos_wc = f"{data.get('rooms', {}).get('bedrooms', '')} + {data.get('rooms', {}).get('bathrooms', '')}"
        conservacao = data.get("visual_analysis", {}).get("estado_conservacao", "")
        indice_fiscal = ""  # Não extraído
        pavimentacao = data.get("visual_analysis", {}).get("pavimentacao", "")
        idade = data.get("visual_analysis", {}).get("idade_aparente_anos", "")
        suites = data.get("rooms", {}).get("suites", "")
        vagas = data.get("rooms", {}).get("parking_spaces", "")
        data_evento = datetime.now().strftime("%d/%m/%Y")
        evento = "Venda"
        geminada = ""  # Não extraído
        multi = ""  # Não extraído
        renda_bairro = ""  # Não extraído
        valor_total = data.get("prices", {}).get("price_formatted", "")

        values = [
            endereco,
            bairro,
            telefone,
            area_terreno,
            area_construida,
            padrao,
            quartos_wc,
            conservacao,
            indice_fiscal,
            pavimentacao,
            idade,
            suites,
            vagas,
            data_evento,
            evento,
            geminada,
            multi,
            renda_bairro,
            valor_total,
        ]

        for col, value in enumerate(values, 1):
            ws.cell(row=row, column=col, value=value)

        # Ajustar largura
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[
                chr(64 + col) if col <= 26 else "A" + chr(64 + col - 26)
            ].width = 15

        wb.save(output_file)
        return True

    except Exception as e:
        print(f"    ⚠️ Erro ao exportar Excel: {e}")
        return False


async def process_batch(urls_file: str, output_base_dir: str = ".tmp/batch_10_imoveis"):
    """
    Processa batch de imóveis
    """

    # Ler URLs
    with open(urls_file, "r") as f:
        urls = [line.strip() for line in f if line.strip()]

    print(f"🎯 Processando {len(urls)} imóveis...")
    print(f"📁 Saída: {output_base_dir}")
    print()

    base_dir = Path(output_base_dir)
    base_dir.mkdir(parents=True, exist_ok=True)

    results = []

    for i, url in enumerate(urls, 1):
        print(f"\n{'=' * 80}")
        print(f"🏠 Imóvel {i}/{len(urls)}")
        print(f"🔗 {url}")
        print("=" * 80)

        property_dir = base_dir / f"property_{i:03d}"

        try:
            # 1. Scraping
            print(f"\n📥 [1/3] Scraping...")
            data = await scrape_single_property(url, property_dir)

            if not data:
                print(f"❌ Falha no scraping")
                results.append(
                    {
                        "index": i,
                        "url": url,
                        "status": "error",
                        "error": "Scraping failed",
                    }
                )
                continue

            print(
                f"✅ Scraping OK: {data.get('basic_info', {}).get('title', 'Sem título')[:50]}..."
            )

            # 2. Análise visual
            print(f"\n🤖 [2/3] Análise visual com Groq...")
            image_paths = [img["local_path"] for img in data.get("images", [])[:3]]

            if image_paths:
                analysis = analyze_with_groq(image_paths)
                if analysis:
                    data["visual_analysis"] = analysis
                    print(
                        f"✅ Análise OK: {analysis.get('estado_conservacao', 'N/A')} / {analysis.get('padrao_construtivo', 'N/A')}"
                    )
                else:
                    print(f"⚠️ Análise falhou")
            else:
                print(f"⚠️ Sem imagens para analisar")

            # Salvar JSON atualizado
            with open(property_dir / "data.json", "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)

            # 3. Excel
            print(f"\n📊 [3/3] Exportando Excel...")
            excel_file = property_dir / "planilha.xlsx"
            if export_to_excel_format(data, str(excel_file)):
                print(f"✅ Excel OK: {excel_file}")
            else:
                print(f"⚠️ Excel falhou")

            results.append(
                {
                    "index": i,
                    "url": url,
                    "status": "success",
                    "title": data.get("basic_info", {}).get("title", ""),
                    "price": data.get("prices", {}).get("price_formatted", ""),
                    "neighborhood": data.get("location", {}).get("neighborhood", ""),
                }
            )

        except Exception as e:
            print(f"❌ Erro: {e}")
            results.append({"index": i, "url": url, "status": "error", "error": str(e)})

        if i < len(urls):
            print(f"\n⏳ Aguardando 10s...")
            await asyncio.sleep(10)

    # Resumo
    print(f"\n{'=' * 80}")
    print(f"✅ BATCH CONCLUÍDO")
    print(f"{'=' * 80}")
    success = len([r for r in results if r["status"] == "success"])
    failed = len([r for r in results if r["status"] == "error"])
    print(f"📊 Total: {len(urls)} | ✅ Sucesso: {success} | ❌ Falhas: {failed}")

    # Salvar resumo
    summary_file = base_dir / "summary.json"
    with open(summary_file, "w", encoding="utf-8") as f:
        json.dump(
            {
                "timestamp": datetime.now().isoformat(),
                "total": len(urls),
                "successful": success,
                "failed": failed,
                "results": results,
            },
            f,
            ensure_ascii=False,
            indent=2,
        )

    print(f"📁 Resumo: {summary_file}")

    return results


if __name__ == "__main__":
    urls_file = sys.argv[1] if len(sys.argv) > 1 else ".tmp/urls_10_imoveis.txt"
    output_dir = sys.argv[2] if len(sys.argv) > 2 else ".tmp/batch_10_imoveis"

    asyncio.run(process_batch(urls_file, output_dir))
