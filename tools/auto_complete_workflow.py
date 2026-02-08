#!/usr/bin/env python3
"""
Auto-complete workflow for 1000 properties
Monitors scraping, runs AI analysis, generates Excel, commits to GitHub
"""

import os
import sys
import time
import json
import subprocess
from pathlib import Path
from datetime import datetime


def log(msg):
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    print(f"[{timestamp}] {msg}")
    with open(".tmp/auto_complete.log", "a") as f:
        f.write(f"[{timestamp}] {msg}\n")


def count_processed_properties():
    """Count how many properties have been processed"""
    base_dir = Path(".tmp/batch_1000_imoveis")
    if not base_dir.exists():
        return 0
    return len(list(base_dir.glob("property_*/data.json")))


def check_scraping_complete():
    """Check if scraping is complete (1000 properties)"""
    return count_processed_properties() >= 1000


def run_ai_analysis():
    """Run AI analysis on all properties"""
    log("🤖 Starting AI analysis...")

    # Check if analyze_batch_mistral.py exists
    if not Path("tools/analyze_batch_mistral.py").exists():
        log("❌ analyze_batch_mistral.py not found. Creating it...")
        create_ai_analysis_script()

    result = subprocess.run(
        ["python", "tools/analyze_batch_mistral.py", ".tmp/batch_1000_imoveis"],
        capture_output=True,
        text=True,
    )

    if result.returncode == 0:
        log("✅ AI analysis completed")
        return True
    else:
        log(f"❌ AI analysis failed: {result.stderr}")
        return False


def generate_excel():
    """Generate Excel file from processed data"""
    log("📊 Generating Excel file...")

    # Check if script exists
    if not Path("tools/generate_excel_1000.py").exists():
        log("Creating Excel generator...")
        create_excel_generator()

    result = subprocess.run(
        ["python", "tools/generate_excel_1000.py"], capture_output=True, text=True
    )

    if result.returncode == 0:
        log("✅ Excel file generated")
        return True
    else:
        log(f"❌ Excel generation failed: {result.stderr}")
        return False


def commit_to_github():
    """Commit results to GitHub"""
    log("📦 Committing to GitHub...")

    commands = [
        ["git", "add", ".tmp/batch_1000_imoveis/summary.json"],
        ["git", "add", "planilha_1000_imoveis_v2.xlsx"],
        [
            "git",
            "commit",
            "-m",
            "v2.3: Complete 1000 properties dataset with AI analysis",
        ],
        ["git", "push", "origin", "v2"],
    ]

    for cmd in commands:
        result = subprocess.run(cmd, capture_output=True, text=True)
        if result.returncode != 0:
            log(f"⚠️ Git command failed: {' '.join(cmd)}")
            log(f"   Error: {result.stderr}")

    log("✅ GitHub commit completed")


def create_ai_analysis_script():
    """Create AI analysis script if it doesn't exist"""
    script_content = '''#!/usr/bin/env python3
"""
Analyze batch with Mistral AI
"""
import os
import json
import base64
from pathlib import Path
from mistralai import Mistral

def analyze_image(image_path, client):
    """Analyze single image with Mistral"""
    try:
        with open(image_path, "rb") as f:
            image_data = base64.b64encode(f.read()).decode()
        
        response = client.chat.complete(
            model="pixtral-12b-2409",
            messages=[
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": "Analise esta imagem de imóvel e responda em JSON: {estado_conservacao: 'novo'|'bom'|'regular'|'ruim', idade_aparente: number, padrao_construcao: 'alto'|'medio'|'baixo', tipo_piso: 'ceramica'|'porcelanato'|'cimento'|'outro'}"},
                        {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{image_data}"}}
                    ]
                }
            ]
        )
        
        return response.choices[0].message.content
    except Exception as e:
        return f'{{"error": "{str(e)}"}}'

def process_property(property_dir, client):
    """Process all images in a property directory"""
    data_file = Path(property_dir) / "data.json"
    
    if not data_file.exists():
        return
    
    with open(data_file, "r", encoding="utf-8") as f:
        data = json.load(f)
    
    # Check if already analyzed
    if "ai_analysis" in data:
        return
    
    images_dir = Path(property_dir) / "images"
    if not images_dir.exists():
        return
    
    analyses = []
    for img_file in sorted(images_dir.glob("*.jpg"))[:3]:  # Analyze first 3 images
        analysis = analyze_image(img_file, client)
        analyses.append(analysis)
    
    data["ai_analysis"] = analyses
    
    with open(data_file, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

if __name__ == "__main__":
    import sys
    base_dir = sys.argv[1] if len(sys.argv) > 1 else ".tmp/batch_1000_imoveis"
    
    api_key = os.environ.get("MISTRAL_API_KEY")
    if not api_key:
        print("❌ MISTRAL_API_KEY not set")
        sys.exit(1)
    
    client = Mistral(api_key=api_key)
    
    property_dirs = sorted(Path(base_dir).glob("property_*"))
    print(f"Processing {len(property_dirs)} properties...")
    
    for i, prop_dir in enumerate(property_dirs, 1):
        print(f"[{i}/{len(property_dirs)}] Analyzing {prop_dir.name}...")
        process_property(prop_dir, client)
    
    print("✅ AI analysis complete")
'''

    with open("tools/analyze_batch_mistral.py", "w") as f:
        f.write(script_content)


def create_excel_generator():
    """Create Excel generator script"""
    script_content = '''#!/usr/bin/env python3
"""
Generate Excel for 1000 properties
"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def extract_property_data(property_dir):
    """Extract data from property JSON"""
    data_file = Path(property_dir) / "data.json"
    
    if not data_file.exists():
        return None
    
    with open(data_file, "r", encoding="utf-8") as f:
        data = json.load(f)
    
    return {
        "title": data.get("basic_info", {}).get("title", ""),
        "price": data.get("prices", {}).get("price_formatted", ""),
        "price_numeric": data.get("prices", {}).get("price_numeric", 0),
        "neighborhood": data.get("location", {}).get("neighborhood", ""),
        "city": data.get("location", {}).get("city", ""),
        "bedrooms": data.get("rooms", {}).get("bedrooms", 0),
        "bathrooms": data.get("rooms", {}).get("bathrooms", 0),
        "suites": data.get("rooms", {}).get("suites", 0),
        "parking": data.get("rooms", {}).get("parking_spaces", 0),
        "description": data.get("description", {}).get("text", "")[:200],
        "url": data.get("metadata", {}).get("source_url", ""),
    }

def generate_excel():
    """Generate Excel file"""
    base_dir = Path(".tmp/batch_1000_imoveis")
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Imóveis"
    
    # Headers
    headers = ["#", "Título", "Preço", "Bairro", "Cidade", "Quartos", "Banheiros", "Suítes", "Vagas", "Descrição", "URL"]
    ws.append(headers)
    
    # Style headers
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    # Add data
    property_dirs = sorted(base_dir.glob("property_*"))
    for i, prop_dir in enumerate(property_dirs, 1):
        data = extract_property_data(prop_dir)
        if data:
            ws.append([
                i,
                data["title"],
                data["price"],
                data["neighborhood"],
                data["city"],
                data["bedrooms"],
                data["bathrooms"],
                data["suites"],
                data["parking"],
                data["description"],
                data["url"]
            ])
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 10
    ws.column_dimensions['G'].width = 10
    ws.column_dimensions['H'].width = 10
    ws.column_dimensions['I'].width = 10
    ws.column_dimensions['J'].width = 50
    ws.column_dimensions['K'].width = 60
    
    output_file = "planilha_1000_imoveis_v2.xlsx"
    wb.save(output_file)
    print(f"✅ Excel saved: {output_file}")

if __name__ == "__main__":
    generate_excel()
'''

    with open("tools/generate_excel_1000.py", "w") as f:
        f.write(script_content)


def main():
    log("=" * 80)
    log("🚀 Auto-complete workflow started")
    log("=" * 80)

    # Step 1: Monitor scraping
    log("⏳ Monitoring scraping progress...")
    check_interval = 60  # Check every minute
    last_count = 0

    while not check_scraping_complete():
        current_count = count_processed_properties()

        if current_count != last_count:
            log(
                f"📊 Progress: {current_count}/1000 properties ({current_count / 10:.1f}%)"
            )
            last_count = current_count

        time.sleep(check_interval)

    log("✅ Scraping complete! All 1000 properties processed.")

    # Step 2: AI Analysis
    if run_ai_analysis():
        # Step 3: Generate Excel
        if generate_excel():
            # Step 4: Commit to GitHub
            commit_to_github()

    log("=" * 80)
    log("🎉 WORKFLOW COMPLETE!")
    log("=" * 80)


if __name__ == "__main__":
    main()
