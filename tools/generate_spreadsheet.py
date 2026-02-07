#!/usr/bin/env python3
"""
Gera planilha consolidada no formato do usuário
===============================================
Usa dados extraídos + análise baseada em regras (sem IA)
"""

import json
import re
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def parse_description(description_text: str) -> dict:
    """
    Extrai informações da descrição textual
    """
    info = {
        "quartos": "",
        "banheiros": "",
        "suites": "",
        "vagas": "",
        "area_terreno": "",
        "area_construida": "",
        "pavimentacao": "",
        "conservacao": "",
    }

    text_lower = description_text.lower()

    # Extrair quartos
    match = re.search(r"(\d+)\s*dormitório", text_lower)
    if match:
        info["quartos"] = match.group(1)

    # Extrair suítes
    match = re.search(r"(\d+)\s*suíte", text_lower)
    if match:
        info["suites"] = match.group(1)

    # Extrair vagas
    match = re.search(r"(\d+)\s*vaga", text_lower)
    if match:
        info["vagas"] = match.group(1)

    # Extrair áreas
    match = re.search(r"(\d+)m².*?terreno", text_lower)
    if match:
        info["area_terreno"] = match.group(1) + " m²"

    match = re.search(r"(\d+)m².*?construída", text_lower)
    if match:
        info["area_construida"] = match.group(1) + " m²"

    # Pavimentação
    if "cerâmica" in text_lower or "ceramica" in text_lower:
        info["pavimentacao"] = "Cerâmica"
    elif "porcelanato" in text_lower:
        info["pavimentacao"] = "Porcelanato"
    elif "madeira" in text_lower:
        info["pavimentacao"] = "Madeira"

    # Estado de conservação baseado em palavras-chave
    if any(
        word in text_lower
        for word in ["novo", "moderno", "recém construído", "impecável"]
    ):
        info["conservacao"] = "Novo"
    elif any(word in text_lower for word in ["reforma", "acabamento", "pintura nova"]):
        info["conservacao"] = "Bom"
    else:
        info["conservacao"] = "Regular"

    return info


def generate_spreadsheet(
    batch_dir: str = ".tmp/batch_resultado",
    output_file: str = ".tmp/planilha_10_imoveis.xlsx",
):
    """
    Gera planilha Excel no formato solicitado pelo usuário
    """

    batch_path = Path(batch_dir)

    # Encontrar todos os diretórios de imóveis
    property_dirs = sorted(
        [
            d
            for d in batch_path.iterdir()
            if d.is_dir() and d.name.startswith("property_")
        ]
    )

    print(f"📊 Gerando planilha para {len(property_dirs)} imóveis...")

    # Criar workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Imóveis"

    # Headers conforme planilha do usuário
    headers = [
        "Endereço",
        "Bairro",
        "Telefone do informante",
        "Área do Terreno",
        "Área Construída",
        "Padrão construtivo",
        "Quartos + Wc",
        "Estado de conservação",
        "Indice fiscal",
        "Pavimentação",
        "Idade aparente",
        "Suítes",
        "Vagas de garagem coberta",
        "Data do evento",
        "Evento",
        "Geminada",
        "Multi",
        "Renda Média Bairro",
        "Valor total",
    ]

    # Estilo header
    header_fill = PatternFill(
        start_color="366092", end_color="366092", fill_type="solid"
    )
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Escrever headers
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = thin_border

    # Ajustar altura da linha de header
    ws.row_dimensions[1].height = 30

    # Processar cada imóvel
    row = 2
    for prop_dir in property_dirs:
        json_file = prop_dir / "data.json"
        if not json_file.exists():
            continue

        with open(json_file, "r", encoding="utf-8") as f:
            data = json.load(f)

        # Extrair informações
        desc_text = data.get("description", {}).get("text", "")
        parsed = parse_description(desc_text)

        # Mapear campos
        endereco = data.get("basic_info", {}).get("title", "N/A")
        bairro = data.get("location", {}).get("neighborhood", "N/A")
        telefone = ""  # Não extraído
        area_terreno = parsed.get("area_terreno", "")
        area_construida = parsed.get("area_construida", "")
        padrao = "Médio"  # Default
        quartos_wc = f"{parsed.get('quartos', '')} + {parsed.get('banheiros', '1')}"
        conservacao = parsed.get("conservacao", "Regular")
        indice_fiscal = ""  # Não disponível
        pavimentacao = parsed.get("pavimentacao", "Indeterminado")
        idade = ""  # Não extraído
        suites = parsed.get("suites", "")
        vagas = parsed.get("vagas", "")
        data_evento = datetime.now().strftime("%d/%m/%Y")
        evento = "Venda"
        geminada = "Não"  # Default
        multi = ""  # Não disponível
        renda_bairro = ""  # Não disponível
        valor = data.get("prices", {}).get("price_formatted", "N/A")

        values = [
            endereco,
            bairro,
            telefone,
            area_terreno,
            area_construida,
            padrao,
            quartos_wc,
            conservacao,
            indice_fiscal,
            pavimentacao,
            idade,
            suites,
            vagas,
            data_evento,
            evento,
            geminada,
            multi,
            renda_bairro,
            valor,
        ]

        # Escrever linha
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        row += 1

    # Ajustar largura das colunas
    column_widths = {
        "A": 40,  # Endereço
        "B": 20,  # Bairro
        "C": 20,  # Telefone
        "D": 15,  # Área Terreno
        "E": 15,  # Área Construída
        "F": 15,  # Padrão
        "G": 12,  # Quartos + WC
        "H": 18,  # Conservação
        "I": 12,  # Indice fiscal
        "J": 15,  # Pavimentação
        "K": 12,  # Idade
        "L": 10,  # Suítes
        "M": 18,  # Vagas
        "N": 12,  # Data
        "O": 12,  # Evento
        "P": 12,  # Geminada
        "Q": 10,  # Multi
        "R": 18,  # Renda
        "S": 20,  # Valor
    }

    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Congelar primeira linha
    ws.freeze_panes = "A2"

    # Salvar
    wb.save(output_file)
    print(f"✅ Planilha gerada: {output_file}")
    print(f"📊 Total de imóveis: {row - 2}")

    return output_file


if __name__ == "__main__":
    import sys

    batch_dir = sys.argv[1] if len(sys.argv) > 1 else ".tmp/batch_resultado"
    output = sys.argv[2] if len(sys.argv) > 2 else ".tmp/planilha_10_imoveis.xlsx"

    generate_spreadsheet(batch_dir, output)
