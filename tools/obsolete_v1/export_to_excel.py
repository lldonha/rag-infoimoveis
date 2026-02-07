"""
Export dados do PostgreSQL para Excel

Lê imóveis do banco e gera planilha Excel formatada
"""
import sys
from pathlib import Path
from datetime import datetime
from typing import List, Dict

sys.path.append(str(Path(__file__).parent.parent))

from tools.property_saver import get_recent_properties


def export_to_excel(properties: List[Dict], output_file: str):
    """Exporta dados para Excel com formatação visual"""
    try:
        import pandas as pd
        from openpyxl import load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        # Criar DataFrame
        df = pd.DataFrame(properties)

        # Reordenar colunas para melhor visualização
        columns_order = [
            'title', 'price_brl', 'property_type', 'transaction_type',
            'neighborhood', 'city', 'state',
            'area_total_m2', 'area_built_m2', 'bedrooms', 'bathrooms', 'suites', 'parking_spaces',
            'price_per_m2', 'condominium_fee_brl', 'iptu_annual_brl',
            'description', 'address', 'features',
            'source_url', 'completeness', 'scraped_at'
        ]

        # Filtrar apenas colunas existentes
        columns_order = [col for col in columns_order if col in df.columns]
        df = df[columns_order]

        # Renomear colunas para português
        column_names = {
            'title': 'Título',
            'price_brl': 'Preço (R$)',
            'property_type': 'Tipo',
            'transaction_type': 'Transação',
            'neighborhood': 'Bairro',
            'city': 'Cidade',
            'state': 'Estado',
            'area_total_m2': 'Área Total (m²)',
            'area_built_m2': 'Área Construída (m²)',
            'bedrooms': 'Quartos',
            'bathrooms': 'Banheiros',
            'suites': 'Suítes',
            'parking_spaces': 'Vagas',
            'price_per_m2': 'R$/m²',
            'condominium_fee_brl': 'Condomínio (R$)',
            'iptu_annual_brl': 'IPTU Anual (R$)',
            'description': 'Descrição',
            'address': 'Endereço',
            'features': 'Características',
            'source_url': 'URL',
            'completeness': 'Completude (%)',
            'scraped_at': 'Data do Scraping'
        }
        df.rename(columns=column_names, inplace=True)

        # Converter completeness para porcentagem
        if 'Completude (%)' in df.columns:
            df['Completude (%)'] = pd.to_numeric(df['Completude (%)'], errors='coerce')
            df['Completude (%)'] = (df['Completude (%)'] * 100).round(1)

        # Salvar para Excel
        df.to_excel(output_file, index=False, sheet_name='Imóveis')

        # Aplicar formatação
        wb = load_workbook(output_file)
        ws = wb['Imóveis']

        # Estilos
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Formatar cabeçalho
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border

        # Formatar células de dados
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical='top', wrap_text=True)

        # Ajustar largura das colunas
        column_widths = {
            'A': 50,  # Título
            'B': 15,  # Preço
            'C': 15,  # Tipo
            'D': 12,  # Transação
            'E': 20,  # Bairro
            'F': 15,  # Cidade
            'G': 8,   # Estado
            'H': 15,  # Área Total
            'I': 15,  # Área Construída
            'J': 10,  # Quartos
            'K': 10,  # Banheiros
            'L': 10,  # Suítes
            'M': 10,  # Vagas
            'N': 12,  # R$/m²
            'O': 15,  # Condomínio
            'P': 15,  # IPTU
            'Q': 60,  # Descrição
            'R': 40,  # Endereço
            'S': 40,  # Características
            'T': 50,  # URL
            'U': 12,  # Completude
            'V': 20,  # Data Scraping
        }

        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width

        # Congelar primeira linha
        ws.freeze_panes = 'A2'

        # Salvar
        wb.save(output_file)

        print(f"\n✅ Planilha Excel criada: {output_file}")
        print(f"   📊 Total de registros: {len(df)}")

        return True

    except ImportError as e:
        print(f"\n❌ Erro: Bibliotecas necessárias não instaladas")
        print(f"   Execute: pip install pandas openpyxl")
        return False
    except Exception as e:
        print(f"\n❌ Erro ao criar planilha: {e}")
        import traceback
        traceback.print_exc()
        return False


def main():
    print("\n" + "=" * 80)
    print("📊 EXPORT PARA EXCEL - InfoImóveis")
    print("=" * 80 + "\n")

    # Ler dados do banco
    print("📥 Lendo dados do PostgreSQL...")
    properties = get_recent_properties(limit=100)

    if not properties:
        print("\n⚠️  Nenhum dado encontrado no banco!")
        return False

    print(f"✅ {len(properties)} imóveis encontrados no banco\n")

    # Gerar nome do arquivo
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    output_dir = Path(__file__).parent.parent / '.tmp'
    output_dir.mkdir(exist_ok=True)
    output_file = output_dir / f'imoveis_scraped_{timestamp}.xlsx'

    # Exportar
    print("📝 Gerando planilha Excel...")
    success = export_to_excel(properties, str(output_file))

    if success:
        print(f"\n🎉 SUCESSO! Planilha criada em:")
        print(f"   📂 {output_file.absolute()}")
        print(f"\n   Abra no Excel para visualizar os dados!")
        return True
    else:
        print("\n❌ Falha ao gerar planilha")
        return False


if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1)
